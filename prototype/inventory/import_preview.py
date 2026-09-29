"""Local bridge from an uploaded file to the T04 parsers, then to ImportBatch.

The real-format parsers live in import-spike/.  An upload is copied to a
temporary file, parsed, and the temporary file is deleted.  The parsed lines
are kept in ImportBatch (local prototype database only) so a person can look
at them and confirm; nothing here writes inventory.  Posting is
inventory.import_posting.
"""

from __future__ import annotations

import hashlib
import re
import sys
from collections import Counter
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import BadZipFile

import openpyxl
from pypdf import PdfReader
from pypdf.errors import PdfReadError

IMPORT_SPIKE_DIR = Path(__file__).resolve().parents[2] / "import-spike"
if str(IMPORT_SPIKE_DIR) not in sys.path:
    sys.path.insert(0, str(IMPORT_SPIKE_DIR))

from import_spike.pdf_order import parse_pdf_order_text  # noqa: E402
from import_spike.real_format import (  # noqa: E402
    INSPECTION_HEADERS,
    STOCK_HEADERS,
    parse_inspection_sheet,
    parse_stock_export,
)

from .models import ImportBatch, Owner, Product  # noqa: E402


class PreviewError(ValueError):
    """A file cannot be previewed safely."""


# What a person should do about each blocking problem, in plain words.
ADVICE = {
    "unit_needs_confirmation": "单据上这几行是整箱商品，但数量栏印的是 EA，系统无法确定是按件还是按箱。请在下方逐行选择单位后再入账。",
    "product_unknown": "这个商品还没有在系统里。先导入该货主的库存表（或验货纸）建立商品，再回到这一页，系统会自动重新检查，不用重新上传。",
    "location_missing": "正常库存没有仓位，入账后找不到货。请在原表补上仓位后重新上传。",
    "expiry_ambiguous": "效期写法看不懂，系统不会猜日期。请在原表改成 2027/07/01 这种写法后重新上传。",
    "code_not_text": "商品编码被 Excel 存成了数字，前面的 0 可能已经丢了。请把该列设为文本后重新上传。",
    "code_invalid": "这一行没有可用的商品编号，请在原表核对。",
    "quantity_not_numeric": "数量不是数字，请在原表核对。",
    "quantity_invalid": "数量必须是不带小数的非负整数，请在原表核对。",
    "quantity_unpaired": "PDF 里的数量和商品对不上号，系统不会猜。请人工对照纸质单据，在「出库」页手工接单。",
    "unit_not_confirmed": "单位不是 CS 或 EA，请在原表核对。",
    "warehouse_missing": "缺少仓库名称，请在原表补上后重新上传。",
    "description_unclear": "商品描述跨了多行，可能混入了别的文字，请对照纸质单据。",
}


def advice_for(code: str) -> str:
    return ADVICE.get(code, "请对照原始单据人工核对。")


def detect_kind(path: Path, file_name: str) -> str:
    """Work out the document kind from the file itself; never from a person's guess."""
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        return ImportBatch.Kind.PDF_ORDER
    if suffix != ".xlsx":
        raise PreviewError("只支持 .xlsx 表格和 .pdf 送货单。")
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, KeyError, BadZipFile, ValueError) as exc:
        raise PreviewError(f"这个 Excel 文件打不开：{exc}") from exc
    try:
        active = workbook.active
        first_row = next(active.iter_rows(max_row=1, values_only=True), ())
        if tuple(first_row[:7]) == STOCK_HEADERS:
            return ImportBatch.Kind.STOCK
        sheet = workbook["Data"] if "Data" in workbook.sheetnames else active
        header = next(sheet.iter_rows(min_row=5, max_row=5, max_col=4, values_only=True), ())
        if tuple(header) == INSPECTION_HEADERS:
            return ImportBatch.Kind.INSPECTION
    finally:
        workbook.close()
    raise PreviewError("认不出这是哪种表格：表头既不是库存表，也不是验货纸。请确认选对了文件。")


def default_doc_no(kind: str, file_name: str) -> str:
    stem = Path(file_name).stem
    if kind == ImportBatch.Kind.INSPECTION:
        match = re.search(r"\bSO\d+\b", stem)
        if match:
            return match.group(0)
    return stem[:80]


def _line_dict(line, extra_errors=()) -> dict:
    errors = [{"code": e.code, "message": e.message} for e in line.errors if e.severity == "block"]
    errors += list(extra_errors)
    return {
        "row": line.row_number,
        "code": line.code if isinstance(line.code, str) else (None if line.code is None else str(line.code)),
        "name": line.name if isinstance(line.name, str) else (None if line.name is None else str(line.name)),
        "qty": line.base_quantity,
        "printed_qty": line.quantity,
        "unit": line.base_unit or line.unit,
        "expiry": line.expiry_date.isoformat() if line.expiry_date else None,
        "location": line.location,
        "warehouse": line.warehouse,
        "condition": line.condition,
        "item_unit": line.raw_values.get("item_unit"),
        "seq": line.raw_values.get("seq"),
        "size": line.raw_values.get("size"),
        "errors": errors,
        "postable": not errors,
    }


def refresh_line(line: dict) -> dict:
    line["postable"] = not line["errors"]
    return line


def parse_upload(upload, *, owner_code: str, external_doc_no: str = "") -> ImportBatch:
    """Parse an uploaded file into an unsaved-to-inventory ImportBatch."""
    owner_code = (owner_code or "").strip()
    if not owner_code:
        raise PreviewError("请选择货主：文件里没有写这批货属于哪家客户，需要你告诉系统。")
    if len(owner_code) > 40:
        raise PreviewError("货主代码最多 40 个字符。")
    suffix = Path(upload.name).suffix.lower()
    digest = hashlib.sha256()
    with NamedTemporaryFile(suffix=suffix) as temp:
        for chunk in upload.chunks():
            digest.update(chunk)
            temp.write(chunk)
        temp.flush()
        path = Path(temp.name)
        kind = detect_kind(path, upload.name)
        doc_no = (external_doc_no or "").strip() or default_doc_no(kind, upload.name)
        doc_date = None
        try:
            if kind == ImportBatch.Kind.STOCK:
                parsed = parse_stock_export(path, owner=owner_code, external_doc_no=doc_no)
                batch_errors = []
            elif kind == ImportBatch.Kind.INSPECTION:
                parsed = parse_inspection_sheet(path, owner=owner_code, external_doc_no=doc_no)
                batch_errors = []
            else:
                reader = PdfReader(path)
                text = "\n".join((page.extract_text() or "") for page in reader.pages)
                result = parse_pdf_order_text(text, owner=owner_code)
                parsed = result.lines
                batch_errors = [{"code": e.code, "message": e.message} for e in result.batch_errors]
                doc_no = result.external_doc_no or doc_no
                doc_date = result.doc_date
        except (OSError, ValueError, KeyError, BadZipFile, PdfReadError) as exc:
            raise PreviewError(f"文件读不出来：{exc}") from exc

    lines = [_line_dict(line) for line in parsed]
    if kind == ImportBatch.Kind.PDF_ORDER:
        apply_product_check(lines, owner_code)
    return ImportBatch(
        kind=kind, file_name=Path(upload.name).name[:200], file_sha256=digest.hexdigest(),
        owner_code=owner_code, external_doc_no=doc_no[:80], lines=lines, batch_errors=batch_errors,
        doc_date=doc_date,
    )


def apply_product_check(lines: list[dict], owner_code: str) -> bool:
    """Mark order lines whose product the owner does not have yet.

    Products can be created after a PDF was uploaded (by importing the stock
    sheet), so this is re-run whenever the batch is shown or posted, never
    frozen at upload time.  Returns True when any line changed.
    """
    owner = Owner.objects.filter(code=owner_code).first()
    known = set(Product.objects.filter(owner=owner).values_list("code", flat=True)) if owner else set()
    changed = False
    for line in lines:
        before = [e for e in line["errors"] if e["code"] == "product_unknown"]
        line["errors"] = [e for e in line["errors"] if e["code"] != "product_unknown"]
        if line["code"] and line["code"] not in known:
            line["errors"].append({"code": "product_unknown", "message": f"货主 {owner_code} 还没有商品 {line['code']}"})
        after = [e for e in line["errors"] if e["code"] == "product_unknown"]
        changed |= bool(before) != bool(after)
        refresh_line(line)
    return changed


def recheck(batch: ImportBatch) -> None:
    """Re-evaluate the checks that depend on what is already in the system."""
    if batch.status == ImportBatch.Status.PREVIEW and batch.kind == ImportBatch.Kind.PDF_ORDER:
        if apply_product_check(batch.lines, batch.owner_code) and batch.pk:
            batch.save(update_fields=["lines"])


def problem_summary(batch: ImportBatch) -> list[dict]:
    """Group blocking problems by kind, with a plain-language next step."""
    counts = Counter()
    sample = {}
    for line in batch.lines:
        for error in line["errors"]:
            counts[error["code"]] += 1
            sample.setdefault(error["code"], error["message"])
    rows = [{"code": c, "count": n, "message": sample[c], "advice": advice_for(c)} for c, n in counts.most_common()]
    for error in batch.batch_errors:
        rows.append({"code": error["code"], "count": 1, "message": error["message"], "advice": advice_for(error["code"])})
    return rows
