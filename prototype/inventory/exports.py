"""Excel exports (T10): stock, stock movements, orders. Read-only: nothing here writes the database.

Every text cell is written as text. openpyxl would otherwise turn a value that starts with "=" into a
formula, and file-imported names can start with anything; control characters openpyxl rejects are stripped.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import queries
from .models import Order, StockMovement

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATE_FMT, TIME_FMT = "yyyy-mm-dd", "yyyy-mm-dd hh:mm:ss"


def _cell(ws, row, col, value):
    if isinstance(value, datetime):
        value = timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
        c = ws.cell(row=row, column=col, value=value)
        c.number_format = TIME_FMT
    elif isinstance(value, date):
        c = ws.cell(row=row, column=col, value=value)
        c.number_format = DATE_FMT
    elif isinstance(value, str):
        c = ws.cell(row=row, column=col, value=ILLEGAL_CHARACTERS_RE.sub("", value))
        c.data_type = "s"  # text, never a formula
    else:
        c = ws.cell(row=row, column=col, value=value)
    return c


def _sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for i, h in enumerate(headers, 1):
        _cell(ws, 1, i, h).font = Font(bold=True)
    count = 0
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            _cell(ws, r, i, v)
        count += 1
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(28, len(str(h)) * 2 + 4))
    return count


def _notes(wb, title, lines):
    ws = wb.create_sheet(title)
    for r, text in enumerate(lines, 1):
        _cell(ws, r, 1, text)
    ws.column_dimensions["A"].width = 100


def _bytes(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stamp() -> str:
    return timezone.localtime().strftime("%Y-%m-%d %H:%M:%S")


def _filtered_product_summaries(balances, *, owner, product, lot, expiry, location, condition):
    """Summarize exactly the stock rows selected by the export filters.

    Product-level, not-yet-allocated reservations have no lot/location/condition, so they are included only
    when the filters still describe whole products. Narrow stock-row filters report allocated quantities from
    those rows instead of silently mixing in unrelated locations or lots.
    """
    keys = {(b["owner"], b["product"]) for b in balances}
    narrow = any((lot, expiry, location, condition))
    if not narrow:
        return [s for s in queries.product_summaries(owner=owner, product=product)
                if (s["owner"], s["product"]) in keys]

    grouped = defaultdict(lambda: {"on_hand": 0, "sellable": 0, "reserved": 0})
    labels = {}
    for b in balances:
        key = (b["owner"], b["product"])
        labels[key] = (b["name"], b["unit"])
        grouped[key]["on_hand"] += b["on_hand"]
        grouped[key]["reserved"] += b["allocated"]
        if b["sellable"]:
            grouped[key]["sellable"] += b["on_hand"]

    result = []
    for owner_code, product_code in sorted(grouped):
        totals = grouped[(owner_code, product_code)]
        name, unit = labels[(owner_code, product_code)]
        result.append({
            "owner": owner_code, "product": product_code, "name": name, "unit": unit,
            "on_hand": totals["on_hand"], "sellable": totals["sellable"],
            "not_sellable": totals["on_hand"] - totals["sellable"],
            "reserved": totals["reserved"], "available": totals["sellable"] - totals["reserved"],
        })
    return result


def inventory_workbook(*, owner="", product="", lot="", expiry="", location="", condition="", show_zero=False) -> bytes:
    filters = dict(owner=owner, product=product, lot=lot, expiry=expiry, location=location, condition=condition)
    wb = Workbook()
    balances = queries.balance_rows(show_zero=show_zero, **filters)
    _sheet(wb, "批次货位明细",
           ["货主", "商品编码", "商品名", "内部批次", "批号", "效期", "货位", "状态", "实物数", "已分配未发", "可自由分配", "批次来源"],
           ([b["owner"], b["product"], b["name"], b["lot"], b["external_lot"], b["expiry"] or "未知", b["location"],
             b["condition_label"], b["on_hand"], b["allocated"], b["free"], b["lot_source"]] for b in balances),
           first=True)
    summaries = _filtered_product_summaries(
        balances, owner=owner, product=product, lot=lot, expiry=expiry, location=location, condition=condition,
    )
    _sheet(wb, "商品汇总", ["货主", "商品编码", "商品名", "单位", "实物数", "可售", "不可售", "占用", "可用（可售−占用）"],
           ([s["owner"], s["product"], s["name"], s["unit"], s["on_hand"], s["sellable"], s["not_sellable"],
             s["reserved"], s["available"]] for s in summaries))
    shown = {k: v for k, v in filters.items() if v}
    _notes(wb, "说明", [
        f"导出时间：{_stamp()}（香港时间）。这是导出那一刻的快照，之后库存还会变。",
        f"筛选条件：{'、'.join(f'{k}={v}' for k, v in shown.items()) or '无'}；{'包含' if show_zero else '不含'}实物数为 0 的行。",
        "商品汇总只合计筛选后的明细；按货位、批次、效期或状态筛选时，占用只含这些明细中已分配未发的数量。",
        "实物数 = 仓库里实际有的；已分配未发 = 已选好货位、还没发出的；可自由分配 = 实物数 − 已分配（只对「可用」状态计）。",
        "可用（可售−占用）是 D07 的候选口径，客户尚未确认（Q03），对账时请以客户认可的口径为准。",
        "本文件只读导出，不改变任何库存数字。",
    ])
    return _bytes(wb)


def movements_workbook() -> bytes:
    wb = Workbook()
    qs = (StockMovement.objects.select_related("balance__owner", "balance__product", "balance__lot")
          .order_by("pk").iterator(chunk_size=2000))
    n = _sheet(wb, "库存流水",
               ["流水号", "时间", "业务日期", "类型", "货主", "商品编码", "内部批次", "货位", "状态", "数量（有正负）",
                "来源单据", "来源行", "操作人"],
               ([m.pk, m.created_at, m.business_date, m.get_kind_display(), m.balance.owner.code, m.balance.product.code,
                 f"L{m.balance.lot_id}", m.location_code, m.get_condition_display(), m.qty, m.source_doc, m.source_line,
                 m.actor] for m in qs), first=True)
    _notes(wb, "说明", [
        f"导出时间：{_stamp()}（香港时间）；共 {n} 条流水，按发生顺序排列。",
        "数量为正 = 增加，为负 = 减少。流水只追加、不修改，可用来核对任意时点的库存。",
        "本文件只读导出，不改变任何库存数字。",
    ])
    return _bytes(wb)


def orders_workbook() -> bytes:
    wb = Workbook()
    orders = Order.objects.select_related("owner").prefetch_related("lines__product", "lines__allocations").order_by("pk")

    def rows():
        for o in orders:
            for line in o.lines.all():
                n = queries.line_numbers(line)
                yield [o.number, o.owner.code, o.created_by, o.created_at, line.line_no, line.product.code,
                       line.requested_expiry or "", n["ordered"], n["shipped"], n["allocated_open"], n["unallocated"],
                       n["cancelled"], o.source_ref, line.source_line]

    count = _sheet(wb, "订单明细",
                   ["订单号", "货主", "接单人", "接单时间", "行号", "商品编码", "指定效期", "订购", "已发", "已分配未发",
                    "未选货位", "已取消", "原件", "原件行"], rows(), first=True)
    _notes(wb, "说明", [
        f"导出时间：{_stamp()}（香港时间）；共 {count} 个订单行。每个订单行一行。",
        "订购 = 已发 + 已分配未发 + 未选货位 + 已取消。",
        "本文件只读导出，不改变任何库存数字。",
    ])
    return _bytes(wb)
