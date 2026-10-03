"""T10 Excel exports: right numbers, read-only, text stays text. Synthetic data only."""
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from inventory import domain
from inventory.models import Order, StockBalance, StockMovement, Operation

from .conftest import ACTOR, EXP_1, new_op

XLSX = "spreadsheetml.sheet"


def book(resp):
    assert resp.status_code == 200 and XLSX in resp["Content-Type"]
    return load_workbook(BytesIO(resp.content))


def table(ws):
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


def db_counts():
    return (StockBalance.objects.count(), sum(StockBalance.objects.values_list("on_hand", "allocated").first() or (0, 0)),
            StockMovement.objects.count(), Operation.objects.count(), Order.objects.count())


def test_inventory_export_matches_the_page_numbers(client, world):
    a = world.opening(10, location="A-01")
    world.opening(5, location="B-01", owner="SYN-OWNER-B")
    res = world.accept("SYN-EXP-1", 4)
    domain.allocate_line(operation_id=new_op(), actor=ACTOR, line_id=world.line_id(res), picks=[{"balance_id": a.pk, "qty": 3}])
    wb = book(client.get(reverse("export_inventory")))
    assert wb.sheetnames == ["批次货位明细", "商品汇总", "说明"]
    head, rows = table(wb["批次货位明细"])
    assert head[:2] == ("货主", "商品编码") and "实物数" in head and len(rows) == 2
    mine = next(r for r in rows if r[0] == "SYN-OWNER-A")
    assert mine[1] == "000777"  # leading zeros survive: text, not the number 777
    assert dict(zip(head, mine))["实物数"] == 10 and dict(zip(head, mine))["已分配未发"] == 3
    assert dict(zip(head, mine))["效期"].date() == EXP_1
    shead, srows = table(wb["商品汇总"])
    s = dict(zip(shead, next(r for r in srows if r[0] == "SYN-OWNER-A")))
    assert (s["实物数"], s["占用"], s["可用（可售−占用）"]) == (10, 4, 6)
    notes = "\n".join(str(c.value) for c in wb["说明"]["A"])
    assert "快照" in notes and "Q03" in notes and "不改变任何库存数字" in notes


def test_inventory_export_follows_the_filters_and_zero_rows(client, world):
    world.opening(10, location="A-01")
    world.opening(5, location="B-01", owner="SYN-OWNER-B")
    _, rows = table(book(client.get(reverse("export_inventory"), {"owner": "SYN-OWNER-B"}))["批次货位明细"])
    assert [r[0] for r in rows] == ["SYN-OWNER-B"]
    _, rows = table(book(client.get(reverse("export_inventory"), {"location": "C-01"}))["批次货位明细"])
    assert rows == []  # an empty result still gives a valid file with headers


def test_inventory_export_summary_follows_location_filter(client, world):
    world.opening(10, location="A-01")
    world.opening(5, location="B-01", owner="SYN-OWNER-B")
    wb = book(client.get(reverse("export_inventory"), {"location": "B-01"}))
    _, detail = table(wb["批次货位明细"])
    head, summary = table(wb["商品汇总"])
    assert len(detail) == len(summary) == 1
    row = dict(zip(head, summary[0]))
    assert (row["货主"], row["实物数"], row["可售"], row["占用"], row["可用（可售−占用）"]) == (
        "SYN-OWNER-B", 5, 5, 0, 5,
    )
    wb = book(client.get(reverse("export_inventory"), {"location": "C-01"}))
    assert table(wb["批次货位明细"])[1] == table(wb["商品汇总"])[1] == []


def test_export_is_read_only_and_get_only(client, world):
    a = world.opening(10)
    res = world.accept("SYN-EXP-2", 2)
    domain.allocate_line(operation_id=new_op(), actor=ACTOR, line_id=world.line_id(res), picks=[{"balance_id": a.pk, "qty": 2}])
    before = db_counts()
    for name in ("export_inventory", "export_movements", "export_orders"):
        assert client.get(reverse(name)).status_code == 200
        assert client.post(reverse(name)).status_code == 405
    assert db_counts() == before


def test_text_starting_with_equals_is_not_a_formula_and_bad_characters_are_dropped(client, world):
    domain.create_product(world.owner_a, "000888", name_zh='=HYPERLINK("http://x","点我")')
    domain.create_product(world.owner_a, "000889", name_zh="坏\x00字符")
    for code in ("000888", "000889"):
        domain.post_opening(operation_id=new_op(), actor=ACTOR, owner_code="SYN-OWNER-A", product_code=code,
                            source_ref=f"SYN-{code}", external_lot="", expiry_date=EXP_1, location_code="A-01",
                            condition="AVAILABLE", qty=1)
    wb = book(client.get(reverse("export_inventory")))
    head, rows = table(wb["批次货位明细"])
    names = {r[1]: r[2] for r in rows}
    assert names["000888"] == '=HYPERLINK("http://x","点我")' and names["000889"] == "坏字符"
    cell = next(c for row in wb["批次货位明细"].iter_rows(min_row=2) for c in row if c.value == names["000888"])
    assert cell.data_type == "s"
    raw = BytesIO(client.get(reverse("export_inventory")).content)
    import zipfile
    assert b"<f>" not in zipfile.ZipFile(raw).read("xl/worksheets/sheet1.xml")  # no formula element at all


def test_movements_export_has_every_row_in_order_with_signed_quantities(client, world):
    a = world.opening(10)
    res = world.accept("SYN-EXP-3", 4)
    domain.allocate_line(operation_id=new_op(), actor=ACTOR, line_id=world.line_id(res), picks=[{"balance_id": a.pk, "qty": 4}])
    order = Order.objects.get(number="SYN-EXP-3")
    alloc = order.lines.get().allocations.get()
    domain.ship(operation_id=new_op(), actor=ACTOR, order_id=order.pk, items=[{"allocation_id": alloc.pk, "qty": 3}])
    resp = client.get(reverse("export_movements"))
    assert "filename*=" in resp["Content-Disposition"] and ".xlsx" in resp["Content-Disposition"]
    head, rows = table(book(resp)["库存流水"])
    assert len(rows) == StockMovement.objects.count() == 2
    d = [dict(zip(head, r)) for r in rows]
    assert [x["类型"] for x in d] == ["期初", "发货"] and [x["数量（有正负）"] for x in d] == [10, -3]
    assert d[1]["来源单据"] == "SYN-EXP-3" and d[1]["操作人"] == ACTOR and d[1]["时间"].year >= 2026


def test_orders_export_one_row_per_line_and_the_numbers_add_up(client, world):
    a = world.opening(20)
    res = domain.accept_order(operation_id=new_op(), actor=ACTOR, owner_code="SYN-OWNER-A", number="SYN-EXP-4",
                              lines=[{"product": "000777", "qty": q} for q in (5, 6, 7)])
    lines = list(Order.objects.get(number="SYN-EXP-4").lines.order_by("line_no"))
    for line, q in zip(lines[:2], (5, 6)):
        domain.allocate_line(operation_id=new_op(), actor=ACTOR, line_id=line.pk, picks=[{"balance_id": a.pk, "qty": q}])
    order = lines[0].order
    alloc = lines[0].allocations.get()
    domain.ship(operation_id=new_op(), actor=ACTOR, order_id=order.pk, items=[{"allocation_id": alloc.pk, "qty": 2}])
    domain.cancel_line(operation_id=new_op(), actor=ACTOR, line_id=lines[1].pk, qty=1)
    head, rows = table(book(client.get(reverse("export_orders")))["订单明细"])
    d = [dict(zip(head, r)) for r in rows]
    assert [x["行号"] for x in d] == [1, 2, 3] and all(x["订单号"] == "SYN-EXP-4" for x in d)
    assert [(x["订购"], x["已发"], x["已分配未发"], x["未选货位"], x["已取消"]) for x in d] == [
        (5, 2, 3, 0, 0), (6, 0, 5, 0, 1), (7, 0, 0, 7, 0)]
    assert all(x["订购"] == x["已发"] + x["已分配未发"] + x["未选货位"] + x["已取消"] for x in d)


def test_empty_database_still_exports_valid_files(client, db):
    for name, sheet in (("export_inventory", "批次货位明细"), ("export_movements", "库存流水"), ("export_orders", "订单明细")):
        head, rows = table(book(client.get(reverse(name)))[sheet])
        assert head and rows == []


def test_pages_have_the_download_buttons(client, world):
    world.opening(3)
    assert reverse("export_inventory") in client.get(reverse("inventory"), {"owner": "SYN-OWNER-A", "show_zero": "1"}).content.decode()
    html = client.get(reverse("inventory"), {"owner": "SYN-OWNER-A", "show_zero": "1"}).content.decode()
    assert "export/inventory/?owner=SYN-OWNER-A" in html and "show_zero=1" in html
    assert reverse("export_orders") in client.get(reverse("outbound")).content.decode()
    assert reverse("export_movements") in client.get(reverse("history")).content.decode()
