"""T04 spike: fixed-template import parsing, validation, preview and duplicate/revision detection.

Scope boundary (see docs/import-contract.md): this module parses a fixed Excel
template into ImportLine-shaped records, flags row-level problems, and classifies
a batch as new / exact_duplicate / revision against a synthetic ledger. It does
NOT write any StockMovement or touch prototype/'s domain layer -- posting a
confirmed batch into real stock is T06's job, calling domain.py's PostImport
command inside a real database transaction (see docs/domain-model.md #5, #7).

Every date/unit/code rule here mirrors an existing, cited requirement:
- D04 (docs/DECISIONS.md): product codes are text, leading zeros preserved;
  expiry is a plain date, not invented from other fields.
- D05 (docs/DECISIONS.md): import is a business operation -- fingerprint the
  file, keep mapping version and row numbers, preview has no side effects,
  same file is a candidate duplicate, different file for the same external
  document is a revision (never silently merged), a note overriding the
  stated quantity needs a human decision, not an automatic override.
- AGENTS.md "库存正确性": never silently merge duplicate-looking rows, infer
  missing dates, convert unknown units, or clamp invalid quantities.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

TEMPLATE_VERSION = "receiving-v1"

# Fixed template contract: exact header text (Chinese) -> internal field name.
# docs/import-contract.md documents this table for the business side.
COLUMN_HEADERS: dict[str, str] = {
    "货主": "owner",
    "外部单号": "external_doc_no",
    "单据版本": "doc_version",
    "商品编码": "code",
    "商品名称": "name",
    "数量": "quantity",
    "单位": "unit",
    "效期": "expiry_raw",
    "外部批号": "external_lot",
    "备注": "remark",
}

REQUIRED_FIELDS = ["owner", "external_doc_no", "code", "name", "quantity", "unit"]

_DATE_RE_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATE_RE_SLASH = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
_REMARK_QTY_RE = re.compile(r"实际数量[:：]?\s*(\d+(?:\.\d+)?)")


@dataclass
class RowError:
    row_number: int
    field: str
    code: str
    message: str
    severity: str  # "block" (cannot auto-post until a human decides) | "warn" (informational)


@dataclass
class ImportLine:
    row_number: int
    owner: str | None
    external_doc_no: str | None
    doc_version: str | None
    code: Any
    name: str | None
    quantity: float | None
    quantity_raw: Any
    unit: str | None
    expiry_raw: Any
    expiry_date: date | None
    expiry_status: str  # "known" | "unknown" | "ambiguous"
    external_lot: str | None
    remark: str | None
    errors: list[RowError] = field(default_factory=list)

    @property
    def postable(self) -> bool:
        return not any(e.severity == "block" for e in self.errors)


@dataclass
class BatchResult:
    file_path: str
    fingerprint: str
    owner: str | None
    external_doc_no: str | None
    doc_version: str | None
    template_version: str
    classification: str  # "new" | "exact_duplicate" | "revision"
    duplicate_of_fingerprint: str | None
    lines: list[ImportLine]
    batch_errors: list[RowError]

    @property
    def postable(self) -> bool:
        if self.classification != "new":
            return False
        if self.batch_errors:
            return False
        return bool(self.lines) and all(line.postable for line in self.lines)


def file_fingerprint(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _read_rows(path: Path, sheet_name: str = "到货") -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    all_rows = list(ws.iter_rows(values_only=False))
    if not all_rows:
        return []
    headers: list[str | None] = []
    for cell in all_rows[0]:
        raw = cell.value
        headers.append(COLUMN_HEADERS.get(str(raw).strip()) if raw is not None else None)
    rows: list[dict] = []
    for row_number, row in enumerate(all_rows[1:], start=2):
        values: dict[str, Any] = {}
        types: dict[str, str] = {}
        for header, cell in zip(headers, row):
            if header is None:
                continue
            values[header] = cell.value
            types[header] = cell.data_type
        if all(v in (None, "") for v in values.values()):
            continue  # trailing blank row
        rows.append({"row_number": row_number, "values": values, "types": types})
    return rows


def _parse_expiry(raw: Any) -> tuple[date | None, str, list[str]]:
    """Returns (best_effort_date_or_None, status, human_readable_notes).

    Only an ISO date (YYYY-MM-DD) or a real date cell counts as unambiguous.
    A blank cell is "unknown" -- not invented. Anything else is "ambiguous":
    we compute a best-effort value only when the day/month order is forced
    by one part being > 12, but the row is still flagged for a human to
    confirm, per D05 -- never silently guessed and posted.
    """
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, "unknown", []
    if isinstance(raw, date):
        return raw, "known", []
    text = str(raw).strip()
    if _DATE_RE_ISO.match(text):
        y, m, d = (int(p) for p in text.split("-"))
        try:
            return date(y, m, d), "known", []
        except ValueError:
            return None, "ambiguous", [f"'{text}' 不是合法日期"]
    m = _DATE_RE_SLASH.match(text)
    if m:
        a, b, y = (int(x) for x in m.groups())
        if a <= 12 and b <= 12 and a != b:
            return None, "ambiguous", [
                f"'{text}' 日/月次序不明确，模板要求 ISO 日期（YYYY-MM-DD），不猜测"
            ]
        day, month = (a, b) if a > 12 else (b, a)
        try:
            guess = date(y, month, day)
        except ValueError:
            return None, "ambiguous", [f"'{text}' 不符合模板要求的 ISO 日期格式"]
        return guess, "ambiguous", [
            f"'{text}' 不是模板要求的 ISO 日期，按数值推出 {guess.isoformat()}，仍需人工确认"
        ]
    return None, "ambiguous", [f"无法识别的日期格式：'{text}'"]


def _normalize_row(
    raw_row: dict,
    unit_conversions: dict,
    owner_hint: str | None,
    doc_no_hint: str | None,
    doc_version_hint: str | None,
) -> ImportLine:
    row_number = raw_row["row_number"]
    values = raw_row["values"]
    types = raw_row["types"]
    errors: list[RowError] = []

    for field_name in REQUIRED_FIELDS:
        if values.get(field_name) in (None, ""):
            errors.append(
                RowError(row_number, field_name, "missing_required", f"缺少必填字段 {field_name}", "block")
            )

    owner = values.get("owner") or owner_hint
    external_doc_no = values.get("external_doc_no") or doc_no_hint
    doc_version = values.get("doc_version") or doc_version_hint

    code = values.get("code")
    if code is not None and types.get("code") != "s":
        errors.append(
            RowError(
                row_number,
                "code",
                "code_not_text",
                f"商品编码被存成非文本类型（{type(code).__name__}: {code!r}），前导零可能已丢失，需要人工核对原始文件",
                "block",
            )
        )
    code_out = code

    quantity_raw = values.get("quantity")
    quantity: float | None
    if isinstance(quantity_raw, (int, float)):
        quantity = float(quantity_raw)
    else:
        quantity = None
        if quantity_raw not in (None, ""):
            errors.append(
                RowError(row_number, "quantity", "quantity_not_numeric", f"数量不是数字：{quantity_raw!r}", "block")
            )
    if quantity is not None and quantity < 0:
        errors.append(RowError(row_number, "quantity", "quantity_negative", f"数量不能为负：{quantity}", "block"))

    remark = values.get("remark") or ""
    remark_match = _REMARK_QTY_RE.search(str(remark))
    if remark_match and quantity is not None:
        remark_qty = float(remark_match.group(1))
        if remark_qty != quantity:
            errors.append(
                RowError(
                    row_number,
                    "quantity",
                    "remark_override_conflict",
                    f"备注写着实际数量 {remark_qty}，与正文数量 {quantity} 不一致，不自动采用备注值，需人工确认",
                    "block",
                )
            )

    unit = values.get("unit")
    if unit is not None:
        base_units = unit_conversions.get("base_units", {})
        confirmed = unit_conversions.get("confirmed", [])
        is_base = bool(base_units.get(unit))
        is_confirmed = any(
            c.get("owner") == owner and c.get("code") == code_out and c.get("unit") == unit for c in confirmed
        )
        if not is_base and not is_confirmed:
            errors.append(
                RowError(
                    row_number,
                    "unit",
                    "unit_not_confirmed",
                    f"单位 '{unit}' 没有该商品已确认的换算版本，不能过账",
                    "block",
                )
            )

    expiry_raw = values.get("expiry_raw")
    expiry_date, expiry_status, notes = _parse_expiry(expiry_raw)
    for note in notes:
        errors.append(RowError(row_number, "expiry_raw", "expiry_ambiguous", note, "block"))

    return ImportLine(
        row_number=row_number,
        owner=owner,
        external_doc_no=external_doc_no,
        doc_version=doc_version,
        code=code_out,
        name=values.get("name"),
        quantity=quantity,
        quantity_raw=quantity_raw,
        unit=unit,
        expiry_raw=expiry_raw,
        expiry_date=expiry_date,
        expiry_status=expiry_status,
        external_lot=values.get("external_lot"),
        remark=values.get("remark"),
        errors=errors,
    )


def _flag_duplicate_rows(lines: list[ImportLine]) -> None:
    """D05: rows that look identical may be legitimate; flag, never merge."""
    seen: dict[tuple, list[ImportLine]] = {}
    for line in lines:
        key = (line.owner, line.code, line.external_lot, str(line.expiry_raw), line.quantity)
        seen.setdefault(key, []).append(line)
    for group in seen.values():
        if len(group) > 1:
            row_numbers = [g.row_number for g in group]
            for line in group:
                line.errors.append(
                    RowError(
                        line.row_number,
                        "*",
                        "duplicate_looking_rows",
                        f"与第 {row_numbers} 行内容相同，可能是合法的多行，未自动合并，请人工确认",
                        "warn",
                    )
                )


def classify_batch(
    fingerprint: str, owner: str | None, external_doc_no: str | None, ledger_batches: list[dict]
) -> tuple[str, str | None]:
    """Classify against a synthetic ledger of previously accepted batches.

    Exact byte-for-byte re-upload -> "exact_duplicate" (A05: retry is a no-op).
    Same (owner, external_doc_no) but different bytes -> "revision" (D05:
    never silently supersede; a human decides which one is authoritative).
    Anything else -> "new".
    """
    for entry in ledger_batches:
        if entry["fingerprint"] == fingerprint:
            return "exact_duplicate", entry["fingerprint"]
    for entry in ledger_batches:
        if entry["owner"] == owner and entry["external_doc_no"] == external_doc_no:
            return "revision", entry["fingerprint"]
    return "new", None


def preview(
    path: Path,
    ledger_batches: list[dict],
    unit_conversions: dict,
    sheet_name: str = "到货",
) -> BatchResult:
    """Pure function: parses and classifies, writes nothing (I9: 预览不产生流水).

    Calling this any number of times with the same ledger must return the
    same classification -- only commit_batch() below is allowed to change
    what counts as "already imported".
    """
    path = Path(path)
    fingerprint = file_fingerprint(path)
    raw_rows = _read_rows(path, sheet_name=sheet_name)

    batch_errors: list[RowError] = []
    if not raw_rows:
        batch_errors.append(RowError(0, "*", "empty_batch", "文件没有数据行", "block"))

    owner_hint = raw_rows[0]["values"].get("owner") if raw_rows else None
    doc_no_hint = raw_rows[0]["values"].get("external_doc_no") if raw_rows else None
    doc_version_hint = raw_rows[0]["values"].get("doc_version") if raw_rows else None

    lines = [
        _normalize_row(row, unit_conversions, owner_hint, doc_no_hint, doc_version_hint) for row in raw_rows
    ]
    _flag_duplicate_rows(lines)

    owners = {ln.owner for ln in lines if ln.owner}
    doc_nos = {ln.external_doc_no for ln in lines if ln.external_doc_no}
    if len(owners) > 1:
        batch_errors.append(
            RowError(0, "owner", "mixed_owner", f"同一批文件出现多个货主：{sorted(owners)}", "block")
        )
    if len(doc_nos) > 1:
        batch_errors.append(
            RowError(0, "external_doc_no", "mixed_doc_no", f"同一批文件出现多个外部单号：{sorted(doc_nos)}", "block")
        )

    owner = next(iter(owners), None)
    external_doc_no = next(iter(doc_nos), None)
    doc_version = lines[0].doc_version if lines else None

    classification, duplicate_of = classify_batch(fingerprint, owner, external_doc_no, ledger_batches)

    return BatchResult(
        file_path=str(path),
        fingerprint=fingerprint,
        owner=owner,
        external_doc_no=external_doc_no,
        doc_version=doc_version,
        template_version=TEMPLATE_VERSION,
        classification=classification,
        duplicate_of_fingerprint=duplicate_of,
        lines=lines,
        batch_errors=batch_errors,
    )


def commit_batch(result: BatchResult, ledger_batches: list[dict], accepted_at: str) -> list[dict]:
    """Simulates accepting a preview into the ledger.

    Returns a NEW list (does not mutate the input) so a caller can tell
    preview() and commit_batch() apart by whether the ledger changed.
    Real posting (StockMovement rows, Operation idempotency row) happens in
    T06's domain layer inside one database transaction; this spike only
    proves the parsing/validation/dedup contract in isolation.
    """
    if not result.postable:
        raise ValueError(
            f"批次不可过账：classification={result.classification}, "
            f"batch_errors={result.batch_errors}, "
            f"blocking_rows={[ln.row_number for ln in result.lines if not ln.postable]}"
        )
    return [
        *ledger_batches,
        {
            "fingerprint": result.fingerprint,
            "owner": result.owner,
            "external_doc_no": result.external_doc_no,
            "doc_version": result.doc_version,
            "accepted_at": accepted_at,
        },
    ]
