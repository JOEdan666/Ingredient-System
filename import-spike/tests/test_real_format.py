from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from import_spike.importer import commit_batch, preview
from import_spike.real_format import convert_to_ea, parse_inspection_sheet, parse_stock_export

ROOT = Path(__file__).resolve().parents[2]
SAMPLES = ROOT / "fixtures" / "synthetic" / "real-format"


def test_stock_export_keeps_rows_and_parses_ng_and_slash_date():
    lines = parse_stock_export(SAMPLES / "stock_export_synthetic.xlsx", owner="SYN-OWNER", external_doc_no="SYN-SNAPSHOT")
    assert len(lines) == 7
    assert [line.row_number for line in lines if line.code == "900103"] == [4, 5, 6, 7]
    assert lines[1].expiry_date == date(2027, 3, 31)
    assert lines[2].condition == "HOLD" and lines[2].location is None
    assert lines[6].code == "009104"
    assert all(line.postable for line in lines)


def test_inspection_unfilled_fields_are_none_and_datetime_is_date():
    lines = parse_inspection_sheet(SAMPLES / "inspection_sheet_synthetic.xlsx", owner="SYN-OWNER", external_doc_no="SYN-INSPECTION")
    assert len(lines) == 3
    assert [line.row_number for line in lines] == [6, 7, 8]
    assert lines[0].raw_values["received"] is None
    assert lines[0].raw_values["over"] is None
    assert lines[0].raw_values["short"] is None
    assert lines[0].expiry_date == date(2027, 11, 27)
    assert type(lines[0].expiry_date) is date
    assert lines[0].base_quantity == 250  # header says 件; do not multiply by 5
    assert convert_to_ea(2, "CS", 5) == 10


def test_case_conversion_rejects_bool_and_fraction():
    import pytest
    with pytest.raises(ValueError):
        convert_to_ea(True, "CS", 5)
    with pytest.raises(ValueError):
        convert_to_ea(1.5, "CS", 5)


def test_inspection_summary_row_is_not_a_product(tmp_path):
    workbook = load_workbook(SAMPLES / "inspection_sheet_synthetic.xlsx")
    workbook["Data"]["A10"] = "SYN TOTAL"
    workbook["Data"]["F10"] = 850
    path = tmp_path / "inspection-summary.xlsx"
    workbook.save(path)
    lines = parse_inspection_sheet(path, owner="SYN-OWNER", external_doc_no="SYN-INSPECTION")
    assert len(lines) == 3


def test_stock_bool_quantity_is_blocked(tmp_path):
    workbook = load_workbook(SAMPLES / "stock_export_synthetic.xlsx")
    workbook.active["G2"] = True
    path = tmp_path / "bad.xlsx"
    workbook.save(path)
    lines = parse_stock_export(path, owner="SYN-OWNER", external_doc_no="SYN-SNAPSHOT")
    assert lines[0].quantity is None
    assert any(error.code == "quantity_not_numeric" for error in lines[0].errors)


def test_existing_preview_rechecks_duplicate_when_committing_stale_preview():
    sample = ROOT / "fixtures" / "synthetic" / "import-samples" / "receiving_normal.xlsx"
    result = preview(sample, [], {"base_units": {"EA": True}})
    first = commit_batch(result, [], "2026-09-29T00:00:00Z")
    import pytest
    with pytest.raises(ValueError):
        commit_batch(result, first, "2026-09-29T00:01:00Z")


def test_old_template_bool_remark_and_version_conflict(tmp_path):
    workbook = load_workbook(ROOT / "fixtures" / "synthetic" / "import-samples" / "receiving_normal.xlsx")
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["数量"], True)
    sheet.cell(3, headers["备注"], "实际数量=-7")
    sheet.cell(3, headers["单据版本"], "2")
    path = tmp_path / "bad-old-template.xlsx"
    workbook.save(path)
    result = preview(path, [], {"base_units": {"EA": True}})
    assert any(error.code == "quantity_not_numeric" for error in result.lines[0].errors)
    assert any(error.code == "remark_override_conflict" for error in result.lines[1].errors)
    assert any(error.code == "mixed_doc_version" for error in result.batch_errors)
