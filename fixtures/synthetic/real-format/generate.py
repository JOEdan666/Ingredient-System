"""Synthetic fixtures that copy the STRUCTURE of the customer's real files, never their values.

Structure observed 2026-09-29 from the real files (kept outside this repo, never committed):
column names, header layout, date spelling, the " NG" warehouse suffix, CS/EA units.
Every code, name, quantity, date and party below is invented.

Run:  cd import-spike && uv run python ../fixtures/synthetic/real-format/generate.py
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent


def stock_export():
    """Mirrors the stock export: one sheet, one header row, 7 named columns plus 2 trailing empty ones (as in the real export).

    Traps copied from the real layout: expiry is TEXT like 2027/3/31 (no zero padding);
    NG stock is marked by the warehouse name suffix " NG" and has an empty location;
    the same code appears in several rows (different expiry / warehouse); no owner column.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["物料编码", "产品中文描述Product Description-Chinese", "规格型号", "仓库名称", "仓位", "有效期至", "库存量(主单位)", None, None])
    rows = [
        ["900101", "SYN 合成犬粮 成犬", "3kg", "SYN-WH01", "AD-01", "2026/11/25", 10],
        ["900102", "SYN 合成犬小食 ", "0.15kg", "SYN-WH01", "AD-01", "2027/3/31", 123],
        ["900103", "SYN 合成猫粮 绝育", "1.5kg", "SYN-WH04 NG", "", "2027/7/16", 2],
        ["900103", "SYN 合成猫粮 绝育", "1.5kg", "SYN-WH04 NG", "", "2027/5/28", 1],
        ["900103", "SYN 合成猫粮 绝育", "1.5kg", "SYN-WH01", "AD-01", "2027/11/11", 120],
        ["900103", "SYN 合成猫粮 绝育", "1.5kg", "SYN-WH01", "AD-02", "2027/7/16", 36],
        ["009104", "SYN 合成单一蛋白 深海", "3kg", "SYN-WH01", "AD-01", "2027/6/30", 3],
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, max_col=1):  # codes are text cells, as in the real export
        for c in row:
            c.number_format = "@"
    wb.save(HERE / "stock_export_synthetic.xlsx")


def inspection_sheet():
    """Mirrors the inspection (驗貨) sheet: title rows, a TWO-row header, blank columns to fill in.

    Traps: item numbers are numeric cells; UOM mixes CS and EA with Piece Per Case;
    expiry is a real datetime; the receiving/over/short/damage columns are empty for staff to fill;
    a second empty worksheet exists.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["SYN SO00001 /"])
    ws.merge_cells("A1:E1")
    ws.append(["（合成）某日拆櫃並安排送貨"])
    ws.append([])
    ws.append(["", "", "", "", "", "存倉數量 (件)", "EXPIRY DATE", "收貨數量", "", "多收", "少收", "破包問題", "效期問題", "封口問題"])
    ws.append(["Item Number", "Description", "Item UOM", "Piece Per Case", "", "", "有效期", "件", "", "件", "件", "件", "件", "件"])
    ws.append([9155801, "SYN Cat Kibble A", "CS", 5, "", 250, datetime(2027, 11, 27)])
    ws.append([9165802, "SYN Cat Kibble B", "EA", 1, "", 440, datetime(2027, 11, 17)])
    ws.append([9165803, "SYN Cat Kibble C", "EA", 1, "", 160, datetime(2027, 12, 15)])
    wb.create_sheet("工作表1")
    wb.save(HERE / "inspection_sheet_synthetic.xlsx")


if __name__ == "__main__":
    stock_export()
    inspection_sheet()
    print("wrote stock_export_synthetic.xlsx, inspection_sheet_synthetic.xlsx")
